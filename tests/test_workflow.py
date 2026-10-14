"""Consumer-visible checks for time-series training, reload, inference, and exports."""

from __future__ import annotations

import math
import tempfile
import unittest
from pathlib import Path

from src.pipeline import analyze_history, forecast, load_training_summary, read_observations, series_keys, train_lstm

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.reporting import export_csv, export_excel


class WorkflowTests(unittest.TestCase):
    @staticmethod
    def observations(periods: int = 124) -> pd.DataFrame:
        dates = pd.date_range("2025-05-01", periods=periods, freq="h")
        index = np.arange(periods)
        pieces = []
        for series, offset in (("series-one", 0.0), ("series-two", 3.2)):
            pieces.append(pd.DataFrame({
                "timestamp": dates,
                "site": series,
                "depth_m": 12.0,
                "temperature_c": 13.0 + offset + 0.012 * index + 0.35 * np.sin(index / 8),
                "dissolved_oxygen_mg_l": 8.5 - 0.007 * index - offset * 0.22 + 0.1 * np.cos(index / 7),
            }))
        return pd.concat(pieces, ignore_index=True).sample(frac=1, random_state=11).reset_index(drop=True)

    @staticmethod
    def _new_observations(directory: Path, loaded: pd.DataFrame) -> Path:
        recent = loaded.groupby("site", sort=False).tail(20)
        new_rows = []
        for _, group in loaded.groupby("site"):
            for offset in range(1, 5):
                row = group.iloc[-1].copy()
                row["timestamp"] += pd.Timedelta(hours=offset)
                row["temperature_c"] += 0.1 * offset
                row["dissolved_oxygen_mg_l"] -= 0.05 * offset
                new_rows.append(row)
        path = directory / "new_observations.xlsx"
        pd.concat([recent, pd.DataFrame(new_rows)], ignore_index=True).to_excel(path, index=False)
        return path

    def test_import_train_reload_predict_and_export(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            csv_path = directory / "history.csv"
            self.observations().to_csv(csv_path, index=False, encoding="utf-8-sig")
            loaded = read_observations(csv_path)
            self.assertEqual(series_keys(loaded), [("series-one", 12.0), ("series-two", 12.0)])
            summary = train_lstm(loaded, directory / "trained", lookback=12, horizon=4, epochs=2)
            self.assertEqual(summary, load_training_summary(directory / "trained"))
            self.assertEqual(summary["sample_counts"]["test"], 2 * (124 - int(124 * 0.85) - 4 + 1))
            early = loaded.groupby("site", sort=False).head(int(124 * 0.70))
            np.testing.assert_allclose(summary["normalization"]["mean"],
                                       early[["temperature_c", "dissolved_oxygen_mg_l"]].mean().values)
            for metric in ("mae", "naive_mae"):
                for feature in ("temperature_c", "dissolved_oxygen_mg_l"):
                    self.assertTrue(math.isfinite(summary["metrics"]["test"][metric][feature]))

            new_data = read_observations(self._new_observations(directory, loaded))
            full = forecast(new_data, directory / "trained", "series-one", 12.0)
            isolated = forecast(new_data[new_data["site"] == "series-one"],
                                directory / "trained", "series-one", 12.0)
            pd.testing.assert_frame_equal(full, isolated)
            self.assertEqual(len(full), 4)
            self.assertGreater(full["timestamp"].iloc[0], new_data["timestamp"].max())
            with self.assertRaisesRegex(ValueError, "not present during training"):
                forecast(new_data, directory / "trained", "untrained-series", 12.0)

            replay = analyze_history(new_data, directory / "trained", "series-one", 12.0)
            series = new_data[new_data["site"] == "series-one"]
            self.assertEqual(replay["sample_count"], len(series) - 12 - 4 + 1)
            self.assertEqual(replay["rows"][0]["timestamp"], series["timestamp"].iloc[15].isoformat())
            self.assertIn("非独立验证", replay["disclaimer"])
            # A target 4 steps ahead cannot influence the model's earlier input window.
            modified = new_data.copy()
            target_index = series.index[15]
            modified.loc[target_index, "temperature_c"] += 5
            changed = analyze_history(modified, directory / "trained", "series-one", 12.0)
            self.assertAlmostEqual(replay["rows"][0]["predicted_temperature_c"],
                                   changed["rows"][0]["predicted_temperature_c"])
            self.assertAlmostEqual(changed["rows"][0]["actual_temperature_c"] -
                                   replay["rows"][0]["actual_temperature_c"], 5.0)

            # Any user-provided label is exported as text, never a spreadsheet formula.
            labelled = full.copy()
            labelled["site"] = "=SUM(1,1)"
            csv_output = export_csv(directory / "forecast.csv", labelled)
            self.assertIn("'=SUM(1,1)", csv_output.read_text(encoding="utf-8-sig"))
            book_path = export_excel(directory / "forecast.xlsx", labelled, summary)
            workbook = load_workbook(book_path)
            self.assertEqual(workbook["短时预测"]["B2"].value, "=SUM(1,1)")
            self.assertEqual(workbook["短时预测"]["B2"].data_type, "s")
            self.assertEqual(workbook["短时预测"].max_row, 5)
            self.assertEqual(workbook["模型评估"]["C4"].value,
                             summary["metrics"]["test"]["mae"]["temperature_c"])
            workbook.close()

    def test_missing_and_duplicate_hours_are_not_silently_used(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "observations.csv"
            frame = self.observations(periods=20)
            frame = frame[~((frame["site"] == "series-one") &
                            (frame["timestamp"] == pd.Timestamp("2025-05-01 05:00")))]
            frame.to_csv(path, index=False)
            with self.assertRaisesRegex(ValueError, "Irregular observation cadence"):
                read_observations(path)
            frame = self.observations(periods=20)
            pd.concat([frame, frame.iloc[[0]]], ignore_index=True).to_csv(path, index=False)
            with self.assertRaisesRegex(ValueError, "Duplicate or out-of-order timestamps"):
                read_observations(path)


if __name__ == "__main__":
    unittest.main()
