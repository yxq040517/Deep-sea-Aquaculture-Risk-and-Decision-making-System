"""Exercise local HTTP import, training, historical analysis, prediction, and export."""

from __future__ import annotations

import io
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from src import server

import numpy as np
import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook


class PlatformApiTests(unittest.TestCase):
    @staticmethod
    def observations(count: int = 112) -> pd.DataFrame:
        index = np.arange(count)
        return pd.DataFrame({
            "timestamp": pd.date_range("2025-05-01", periods=count, freq="h"),
            "site": ["sensor-one"] * count,
            "depth_m": [12.0] * count,
            "temperature_c": 12.0 + 0.02 * index + 0.15 * np.sin(index / 6),
            "dissolved_oxygen_mg_l": 8.7 - 0.006 * index + 0.08 * np.cos(index / 5),
        })

    def test_import_train_replay_predict_export_and_reload_catalog(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch.object(server, "DATASETS", root / "datasets"), patch.object(server, "MODELS", root / "models"):
                with TestClient(server.app) as client:
                    self.assertEqual(client.get("/api/datasets").json()["datasets"], [])
                    self.assertEqual(client.get("/api/models").json()["models"], [])
                    source = self.observations()
                    imported = client.post("/api/datasets", files={
                        "file": ("history.csv", source.to_csv(index=False).encode(), "text/csv")
                    })
                    self.assertEqual(imported.status_code, 200, imported.text)
                    dataset = imported.json()
                    self.assertEqual(dataset["rows"], len(source))
                    self.assertEqual(dataset["series"][0]["site"], "sensor-one")
                    self.assertEqual(dataset["quality"]["duplicate_count"], 0)
                    self.assertEqual(dataset["quality"]["missing_count"], 0)
                    self.assertEqual(dataset["provenance"]["source_name"], "未提供")
                    training = client.post("/api/train", json={
                        "dataset_id": dataset["id"], "lookback": 12, "horizon": 4, "epochs": 2,
                    })
                    self.assertEqual(training.status_code, 200, training.text)
                    job_id = training.json()["job_id"]
                    deadline = time.monotonic() + 50
                    while time.monotonic() < deadline:
                        job = client.get(f"/api/train/{job_id}").json()
                        if job["status"] in ("completed", "failed"):
                            break
                        time.sleep(0.1)
                    self.assertEqual(job["status"], "completed", job)
                    model_id = job["model_id"]
                    self.assertTrue((root / "models" / model_id / "model.pt").is_file())
                    self.assertTrue((root / "models" / model_id / "metadata.json").is_file())
                    self.assertTrue(client.get("/api/models").json()["models"])
                    self.assertEqual(job["summary"]["training_source_sha256"], dataset["sha256"])
                    self.assertEqual(job["summary"]["split_mode"], "chronological")

                    selection = {"dataset_id": dataset["id"], "model_id": model_id,
                                 "site": "sensor-one", "depth_m": 12.0}
                    analysis = client.post("/api/analyze", json=selection)
                    self.assertEqual(analysis.status_code, 200, analysis.text)
                    replay = analysis.json()
                    self.assertEqual(replay["sample_count"], len(source) - 12 - 4 + 1)
                    self.assertIn("非独立验证", replay["disclaimer"])
                    self.assertAlmostEqual(replay["rows"][0]["actual_temperature_c"],
                                           source["temperature_c"].iloc[15])
                    self.assertEqual(replay["validation_context"]["status"], "overlap")
                    self.assertIsNone(replay["event_evaluation"])
                    rule = {"max_temp_c": 12.5, "min_do_mg_l": 8.4,
                            "source": "站点管理试验假设；非鱼体耐受值", "action_minutes": 90}
                    scored = client.post("/api/analyze", json={**selection, "rule": rule})
                    self.assertEqual(scored.status_code, 200, scored.text)
                    self.assertEqual(scored.json()["event_evaluation"]["source"], rule["source"])
                    self.assertIn("temperature", scored.json()["event_evaluation"]["by_variable"])

                    recent = source.tail(20)
                    extension = []
                    for offset in range(1, 5):
                        row = source.iloc[-1].copy()
                        row["timestamp"] += pd.Timedelta(hours=offset)
                        row["temperature_c"] += offset * 0.02
                        row["dissolved_oxygen_mg_l"] -= offset * 0.03
                        extension.append(row)
                    new_data = pd.concat([recent, pd.DataFrame(extension)], ignore_index=True)
                    future_upload = client.post("/api/datasets", files={
                        "file": ("new_observations.xlsx", self._excel_bytes(new_data),
                                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    })
                    self.assertEqual(future_upload.status_code, 200, future_upload.text)
                    selection["dataset_id"] = future_upload.json()["id"]
                    predicted = client.post("/api/forecast", json=selection)
                    self.assertEqual(predicted.status_code, 200, predicted.text)
                    result = predicted.json()
                    self.assertEqual(len(result["history"]), 12)
                    self.assertEqual(len(result["forecast"]), 4)
                    self.assertGreater(result["forecast"][0]["timestamp"], result["history"][-1]["timestamp"])
                    self.assertIn("非实测", result["disclaimer"])
                    self.assertIsNone(result["lead_time"])
                    advised = client.post("/api/forecast", json={**selection, "rule": rule})
                    self.assertEqual(advised.status_code, 200, advised.text)
                    self.assertEqual(advised.json()["lead_time"]["source"], rule["source"])
                    self.assertIn(advised.json()["lead_time"]["status"], ("crossing", "no_crossing"))
                    unsourced = client.post("/api/forecast", json={**selection, "rule": {"max_temp_c": 17}})
                    self.assertEqual(unsourced.status_code, 400)
                    for extension_name in ("csv", "xlsx"):
                        download = client.get(f"/api/forecast/export.{extension_name}", params=selection)
                        self.assertEqual(download.status_code, 200, download.text)
                        self.assertTrue(download.content)
                        if extension_name == "xlsx":
                            workbook = load_workbook(io.BytesIO(download.content), read_only=True)
                            self.assertEqual(workbook["短时预测"].max_row, 5)
                            workbook.close()
                    sourced_report = client.get("/api/forecast/export.xlsx", params={**selection, "rule": json.dumps(rule)})
                    self.assertEqual(sourced_report.status_code, 200, sourced_report.text)
                    workbook = load_workbook(io.BytesIO(sourced_report.content), read_only=True)
                    self.assertEqual(workbook["管理限值与提前量"]["B2"].value, rule["source"])
                    self.assertIn("观测数据来源", workbook.sheetnames)
                    workbook.close()
                    invalid = client.post("/api/forecast", json={**selection, "model_id": "../etc/passwd"})
                    self.assertEqual(invalid.status_code, 400)

                # A new browser/session reads the same persisted dataset and model records.
                with TestClient(server.app) as reopened:
                    self.assertEqual(len(reopened.get("/api/datasets").json()["datasets"]), 2)
                    self.assertEqual(reopened.get("/api/models").json()["models"][0]["id"], model_id)

    def test_upload_quality_provenance_and_single_year_rejection(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch.object(server, "DATASETS", root / "datasets"), patch.object(server, "MODELS", root / "models"):
                with TestClient(server.app) as client:
                    data = self.observations(count=60)
                    data["F_Temp"] = ["<0>"] * len(data)
                    data["F_DO_mgl"] = ["<0>"] * (len(data) - 1) + ["<1>"]
                    provenance = {"source_name": "现场记录", "source_reference": "原始测量簿",
                                  "timezone_note": "原文件未标时区，按 UTC 解读",
                                  "depth_basis": "nominal", "qc_note": "源标记保持原样，含人工复核",
                                  "temperature_unit": "°C", "oxygen_unit": "mg/L"}
                    response = client.post("/api/datasets", data={"provenance": json.dumps(provenance)},
                                           files={"file": ("flagged.csv", data.to_csv(index=False).encode(), "text/csv")})
                    self.assertEqual(response.status_code, 200, response.text)
                    uploaded = response.json()
                    self.assertEqual(uploaded["provenance"]["depth_basis"], "nominal")
                    self.assertEqual(uploaded["quality"]["qc_flags"]["F_DO_mgl"]["<1>"], 1)
                    self.assertEqual(uploaded["quality"]["missing_count"], 0)
                    self.assertTrue(any("标称" in note for note in uploaded["quality"]["warnings"]))
                    persisted = client.get("/api/datasets").json()["datasets"][0]
                    self.assertEqual(persisted["provenance"]["source_reference"], "原始测量簿")

                    gapped = data.drop(index=5)
                    failed = client.post("/api/datasets", files={"file": ("gap.csv", gapped.to_csv(index=False).encode(), "text/csv")})
                    self.assertEqual(failed.status_code, 400)
                    gap_quality = failed.json()["detail"]["quality"]
                    self.assertEqual(gap_quality["missing_count"], 1)
                    self.assertEqual(gap_quality["missing_periods"][0]["missing_steps"], 1)
                    self.assertEqual(len(client.get("/api/datasets").json()["datasets"]), 1)

                    percent = data.copy()
                    percent["dissolved_oxygen_mg_l"] = "80%"
                    rejected = client.post("/api/datasets", files={"file": ("percent.csv", percent.to_csv(index=False).encode(), "text/csv")})
                    self.assertEqual(rejected.status_code, 400)
                    self.assertIn("mg/L", rejected.json()["detail"]["message"])

                    started = client.post("/api/train", json={"dataset_id": uploaded["id"],
                                                              "lookback": 12, "horizon": 4,
                                                              "epochs": 1, "split_mode": "year_holdout"})
                    self.assertEqual(started.status_code, 200, started.text)
                    deadline = time.monotonic() + 30
                    while time.monotonic() < deadline:
                        job = client.get("/api/train/" + started.json()["job_id"]).json()
                        if job["status"] in ("completed", "failed"):
                            break
                        time.sleep(0.05)
                    self.assertEqual(job["status"], "failed")
                    self.assertIn("three calendar years", job["error"])
                    self.assertEqual(client.get("/api/models").json()["models"], [])

    @staticmethod
    def _excel_bytes(frame: pd.DataFrame) -> bytes:
        stream = io.BytesIO()
        frame.to_excel(stream, index=False)
        return stream.getvalue()


if __name__ == "__main__":
    unittest.main()
