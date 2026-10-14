"""Export raw short-horizon forecasts and recorded model evaluation results."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_COLUMNS = {
    "timestamp": "预测时间",
    "site": "观测序列 / 场址",
    "depth_m": "观测水深 (m)",
    "temperature_c": "预测水温 (°C)",
    "dissolved_oxygen_mg_l": "预测溶氧 (mg/L)",
}


def _require_forecast(frame: pd.DataFrame) -> None:
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        raise ValueError("请先生成非空的预测结果")
    if not set(_COLUMNS).issubset(frame.columns):
        raise ValueError("预测结果字段不完整，无法导出")


def _csv_site(value: str) -> str:
    """Stop a user-supplied site identifier becoming a spreadsheet formula."""
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def export_csv(path: str | Path, forecast: pd.DataFrame) -> Path:
    """Write only the numeric forecast and its series identifiers, never a risk claim."""
    _require_forecast(forecast)
    destination = Path(path)
    if destination.suffix.lower() != ".csv":
        raise ValueError("导出路径须为 .csv 文件")
    data = forecast.loc[:, list(_COLUMNS)].copy()
    data["site"] = data["site"].map(_csv_site)
    data.to_csv(destination, index=False, encoding="utf-8-sig")
    return destination


def _cell(sheet, row: int, column: int, value) -> None:
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime() if value.tzinfo is None else value.isoformat()
    elif isinstance(value, datetime) and value.tzinfo is not None:
        value = value.isoformat()
    item = sheet.cell(row, column, value)
    # Excel must treat imported labels as literal text even when they start with '='.
    if isinstance(value, str):
        item.data_type = "s"
    item.alignment = Alignment(vertical="center", wrap_text=True)


def _header(sheet, labels: list[str]) -> None:
    for column, label in enumerate(labels, 1):
        _cell(sheet, 1, column, label)
        item = sheet.cell(1, column)
        item.fill = PatternFill("solid", fgColor="17365D")
        item.font = Font(color="FFFFFF", bold=True)
    sheet.row_dimensions[1].height = 27
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def _widths(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def export_excel(path: str | Path, forecast: pd.DataFrame, summary: dict) -> Path:
    """Save forecasts, the trained time scale, and out-of-sample vs naive errors."""
    _require_forecast(forecast)
    if not isinstance(summary, dict) or summary.get("schema_version") != 1:
        raise ValueError("模型评估记录不存在或格式不兼容")
    destination = Path(path)
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("导出路径须为 .xlsx 文件")
    book = Workbook()
    sheet = book.active
    sheet.title = "短时预测"
    fields = list(_COLUMNS)
    _header(sheet, [*(_COLUMNS[key] for key in fields), "数据类型"])
    for index, record in enumerate(forecast.loc[:, fields].itertuples(index=False, name=None), 2):
        for column, value in enumerate(record, 1):
            _cell(sheet, index, column, value)
        _cell(sheet, index, len(fields) + 1, "模型预测，非实测")
    sheet.auto_filter.ref = sheet.dimensions
    _widths(sheet, [23, 27, 18, 22, 24, 24])

    metrics = book.create_sheet("模型评估")
    _header(metrics, ["数据分段", "变量与单位", "LSTM MAE", "末值基线 MAE"])
    row = 2
    for split, title in (("validation", "验证段"), ("test", "独立测试段")):
        measured = summary.get("metrics", {}).get(split, {})
        for feature, label in (("temperature_c", "水温 / ℃"),
                               ("dissolved_oxygen_mg_l", "溶氧 / mg/L")):
            for column, value in enumerate((title, label, measured.get("mae", {}).get(feature),
                                             measured.get("naive_mae", {}).get(feature)), 1):
                _cell(metrics, row, column, value)
            row += 1
    _widths(metrics, [22, 24, 20, 24])

    info = book.create_sheet("模型信息")
    _header(info, ["记录项", "训练时的值"])
    values = [
        ("模型", "LSTM，多步水温与溶氧预测"),
        ("历史输入 / 采样步", summary["lookback"]),
        ("预测时距 / 采样步", summary["horizon"]),
        ("采样间隔", summary["cadence"]),
        ("训练样本数", summary["sample_counts"]["train"]),
        ("验证样本数", summary["sample_counts"]["validation"]),
        ("测试样本数", summary["sample_counts"]["test"]),
        ("适用范围", "本次短时水温与溶氧预测；非全年适养判断"),
    ]
    for index, (label, value) in enumerate(values, 2):
        _cell(info, index, 1, label)
        _cell(info, index, 2, value)
    _widths(info, [34, 73])
    book.save(destination)
    return destination
